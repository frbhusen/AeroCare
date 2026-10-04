"""Excel (.xlsx via openpyxl) and branded PDF (services/pdf.py) rendering of a report envelope
produced by service.run(). Callers must have already run the report through the permission and
scope checks (service.run(..., export=True)); this module only formats what it is given."""
import io
from datetime import date
from decimal import Decimal

from flask import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from backend.app.core.timeutil import local_now
from backend.app.extensions import db
from .common import label

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_TEXT = {
    "period": ("Period", "الفترة"), "generated": ("Generated", "تاريخ الإنشاء"), "center": ("Health center", "المركز"),
    "totals": ("Total", "الإجمالي"), "truncated": ("Only the first {n} rows are included.",
                                                 "تم تضمين أول {n} صفاً فقط."),
    "group_by": ("Grouped by", "مجمع حسب"), "view": ("View", "العرض"), "days": ("Days", "الأيام"),
}


def _t(key, lang, **kw):
    en, ar = _TEXT[key]
    return (ar if lang == "ar" else en).format(**kw)


def _col_label(c, lang):
    return (c.get("label_ar") or c["label"]) if lang == "ar" else c["label"]


def _center(p):
    from backend.app.models import HealthCenter
    return db.session.get(HealthCenter, p.center_id)


def _header_pairs(rep, lang):
    f = rep["filters"]
    pairs = [[_t("period", lang), f"{f['date_from']} - {f['date_to']}"]]
    for key in ("department", "clinic", "doctor"):
        if rep["filter_labels"].get(key):
            pairs.append([label(key, lang), rep["filter_labels"][key]])
    if f.get("group_by"):
        pairs.append([_t("group_by", lang), label(f["group_by"], lang)])
    if f.get("view"):
        pairs.append([_t("view", lang), f["view"]])
    if f.get("days") is not None and rep["filters"].get("view") == "expiry":
        pairs.append([_t("days", lang), str(f["days"])])
    pairs.append([_t("generated", lang), local_now().strftime("%m/%d/%Y %H:%M")])
    return pairs


def filename(rep, ext):
    f = rep["filters"]
    return f"report-{rep['report']}-{f['date_from']}-to-{f['date_to']}.{ext}"


# ---- xlsx ----------------------------------------------------------------------------
def _xl_value(v, type_):
    if v is None or v == "":
        return None
    if type_ == "int":
        try:
            return int(v)
        except (TypeError, ValueError):
            return str(v)
    if type_ in ("money", "number", "percent"):
        try:
            return Decimal(str(v))
        except Exception:
            return str(v)
    if type_ == "date" and isinstance(v, str) and len(v) >= 10:
        try:
            return date.fromisoformat(v[:10])
        except ValueError:
            return v
    return str(v)


def _put(ws, r, c, value, type_="text", bold=False):
    cell = ws.cell(row=r, column=c)
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.data_type = "s"  # never interpret user text as a formula
    else:
        cell.value = value
    if type_ == "money":
        cell.number_format = "#,##0.00"
    elif type_ == "percent":
        cell.number_format = "0.0"
    elif type_ == "date" and isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"
    if bold:
        cell.font = Font(bold=True)
    return cell


def to_xlsx(p, rep, lang="en"):
    wb = Workbook()
    ws = wb.active
    title = rep["title_ar"] if lang == "ar" else rep["title"]
    ws.title = "".join(ch for ch in rep["report"] if ch.isalnum() or ch == "_")[:31] or "report"
    if lang == "ar":
        ws.sheet_view.rightToLeft = True
    center = _center(p)
    r = 1
    _put(ws, r, 1, title).font = Font(bold=True, size=14)
    r += 1
    _put(ws, r, 1, _t("center", lang), bold=True)
    _put(ws, r, 2, center.name if center else "")
    for k, v in _header_pairs(rep, lang):
        r += 1
        _put(ws, r, 1, k, bold=True)
        _put(ws, r, 2, v)
    if rep["meta"].get("currency"):
        r += 1
        _put(ws, r, 1, "Currency" if lang != "ar" else "العملة", bold=True)
        _put(ws, r, 2, rep["meta"]["currency"])
    r += 2
    header_row = r
    cols = rep["columns"]
    fill = PatternFill("solid", fgColor="0F766E")
    for i, c in enumerate(cols, start=1):
        cell = _put(ws, r, i, _col_label(c, lang))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rep["rows"]:
        r += 1
        for i, c in enumerate(cols, start=1):
            _put(ws, r, i, _xl_value(row.get(c["key"]), c["type"]), c["type"])
    totals = rep.get("totals")
    if totals:
        r += 1
        for i, c in enumerate(cols, start=1):
            v = _t("totals", lang) if i == 1 else _xl_value(totals.get(c["key"]), c["type"])
            _put(ws, r, i, v, c["type"], bold=True)
    r += 1  # blank row after the table
    if rep.get("truncated"):
        r += 1
        _put(ws, r, 1, _t("truncated", lang, n=rep["row_limit"]))
    for note in rep.get("notes") or []:
        r += 1
        _put(ws, r, 1, note).font = Font(italic=True, color="666666")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for i, c in enumerate(cols, start=1):
        width = max([len(_col_label(c, lang))] + [len(str(row.get(c["key"]) or "")) for row in rep["rows"][:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(data, name):
    safe = "".join(ch for ch in name if ch.isalnum() or ch in "._-") or "report.xlsx"
    return Response(data, mimetype=XLSX_MIME, headers={"Content-Disposition": f'attachment; filename="{safe}"',
                                                       "Cache-Control": "no-store"})


# ---- pdf -------------------------------------------------------------------------------
def _fmt(v, type_):
    if v is None:
        return ""
    if type_ == "money":
        try:
            return f"{Decimal(str(v)):,.2f}"
        except Exception:
            return str(v)
    return str(v)


def to_pdf(p, rep, lang="en"):
    from backend.app.models import Department
    from backend.app.services import documents
    from backend.app.services.pdf import render_document
    center = _center(p)
    dept = None
    dep_id = rep["filters"].get("department_id")
    if dep_id:
        dept = db.session.execute(select(Department).where(Department.id == dep_id, p.tenant(Department))
                                  ).scalar_one_or_none()
    cols = rep["columns"]
    num = ("int", "money", "number", "percent")
    rows = [[_fmt(row.get(c["key"]), c["type"]) for c in cols] for row in rep["rows"]]
    if rep.get("totals"):
        rows.append([_t("totals", lang) if i == 0 else _fmt(rep["totals"].get(c["key"]), c["type"])
                     for i, c in enumerate(cols)])
    pairs = _header_pairs(rep, lang)
    if rep["meta"].get("currency"):
        pairs.append(["Currency" if lang != "ar" else "العملة", rep["meta"]["currency"]])
    sections = [{"type": "key_values", "items": pairs, "columns": 2},
                {"type": "table", "columns": [_col_label(c, lang) for c in cols], "rows": rows,
                 "widths": [3 if c["type"] == "text" else 1.4 for c in cols],
                 "align": ["end" if c["type"] in num else "start" for c in cols]}]
    if rep.get("truncated"):
        sections.append({"type": "note", "text": _t("truncated", lang, n=rep["row_limit"])})
    for note in rep.get("notes") or []:
        sections.append({"type": "note", "text": note})
    title = rep["title_ar"] if lang == "ar" else rep["title"]
    return render_document(center=center, department=dept, kind="report", title=title, sections=sections,
                           lang="ar" if lang == "ar" else "en", logo_bytes=documents.load_logo(center),
                           subtitle=f"{rep['filters']['date_from']} - {rep['filters']['date_to']}")
