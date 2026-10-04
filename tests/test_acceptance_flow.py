"""End-to-end acceptance flow (spec §109) through the public HTTP API, one realistic day in a center.

Each step uses the role that would really do it. Module tests cover the edge cases; this test proves the
modules work *together*: patient → appointment → visit/records → prescription → pharmacy → lab → billing
→ files → reports, plus offline replay, undo, and isolation checks along the way.
"""
import io
import uuid

from openpyxl import load_workbook

from tests.conftest import ApiClient

V = "/api/v1"
DAY = "2026-11-17"


def ok(r, *codes):
    codes = codes or (200, 201, 202)
    assert r.status_code in codes, (r.status_code, r.get_json())
    return r.get_json()


def png_bytes():
    # Minimal valid 1x1 PNG.
    return bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                         "1f15c4890000000d4944415478da63f8cfc0f01f0005000201ad2a8c0a0000000049454e44ae426082")


def test_full_center_day(app, world, client_for):
    A, B = world["A"], world["B"]
    cl = A["clinics"]
    rec = client_for(A["users"]["rec_dent"])          # dentistry receptionist
    dent = client_for(A["users"]["dent_doc1"])        # dentist, clinic dent1
    gen = client_for(A["users"]["gen_doc"])           # general medicine
    lab = client_for(A["users"]["lab_doc"])
    pharm = client_for(A["users"]["pharm_doc"])
    mgr = client_for(A["users"]["manager"])
    outsider = client_for(B["users"]["manager"])      # another health center

    # 1. Receptionist registers a patient (duplicate lookup first), sequential code.
    assert ok(rec.get(f"{V}/patients/lookup", query_string={"q": "Ali Ahmad"}))["items"] == []
    pt = ok(rec.post(f"{V}/patients", json={"full_name": "Ali Ahmad", "phone": "0944 111 222",
                                             "clinic_id": cl["dent1"], "gender": "male", "allergies": "Penicillin"}))
    pid = pt["id"]
    assert pt["display_code"].startswith("PAT-")
    assert ok(dent.get(f"{V}/patients", query_string={"q": "ali"}))["items"][0]["id"] == pid
    # Search by phone works too; another clinic's doctor can't see the patient.
    assert ok(rec.get(f"{V}/patients", query_string={"q": "0944111"}))["items"][0]["id"] == pid
    assert client_for(A["users"]["dent_doc2"]).get(f"{V}/patients/{pid}").status_code == 404
    assert outsider.get(f"{V}/patients/{pid}").status_code == 404

    # 2. Appointments: booked by reception, conflict prevented, recurring course, WhatsApp link.
    appt = ok(rec.post(f"{V}/appointments", json={"patient_id": pid, "clinic_id": cl["dent1"],
                                                    "doctor_id": A["users"]["dent_doc1"]["id"],
                                                    "starts_at": f"{DAY}T10:00:00", "ends_at": f"{DAY}T10:30:00"}))
    clash = rec.post(f"{V}/appointments", json={"patient_id": pid, "clinic_id": cl["dent1"],
                                                 "starts_at": f"{DAY}T10:15:00", "ends_at": f"{DAY}T10:45:00"})
    assert clash.status_code == 409
    ok(rec.post(f"{V}/appointments/recurring", json={"patient_id": pid, "clinic_id": cl["dent1"],
                                                       "starts_at": "2026-12-01T09:00:00", "duration_minutes": 30,
                                                       "rule": {"freq": "weekly", "count": 3}}))
    wa = ok(rec.get(f"{V}/appointments/{appt['id']}/whatsapp", query_string={"lang": "ar"}))
    assert wa["url"].startswith("https://wa.me/963944111222")
    appt = ok(rec.post(f"{V}/appointments/{appt['id']}/status", json={"status": "arrived", "version": appt["version"]}))

    # 3. Dentist charts the visit: odontogram + treatment (with visit) + prescription.
    ok(dent.put(f"{V}/dentistry/patients/{pid}/odontogram/permanent/14",
                json={"clinic_id": cl["dent1"], "condition": "decay"}), 200, 201)
    tr = ok(dent.post(f"{V}/dentistry/treatments", json={"patient_id": pid, "clinic_id": cl["dent1"],
                                                          "tooth_number": 14, "procedure": "filling",
                                                          "status": "completed", "fee": "150000.00",
                                                          "create_visit": True}))
    rx = ok(dent.post(f"{V}/prescriptions", json={"patient_id": pid, "clinic_id": cl["dent1"],
                                                    "items": [{"medication_name": "Ibuprofen 400mg", "dose": "1 tab",
                                                               "frequency": "3x daily", "quantity": 10}]}))
    pdf = dent.get(f"{V}/prescriptions/{rx['id']}/pdf")
    assert pdf.status_code == 200 and pdf.data[:5] == b"%PDF-"

    # 4. Pharmacy dispenses from stock (inventory receive first).
    item = ok(pharm.post(f"{V}/inventory/items", json={"name": "Ibuprofen 400mg", "is_medication": True,
                                                        "barcode": f"62{uuid.uuid4().int % 10**11:011d}",
                                                        "selling_price": "500.00"}))
    meta = ok(pharm.get(f"{V}/pharmacy/meta"))
    ok(pharm.post(f"{V}/inventory/receive", json={"location_id": meta["clinics"][0]["location_id"], "item_id": item["id"],
                                                   "quantity": 50, "expiry_date": "2027-06-30"}))
    queue = ok(pharm.get(f"{V}/pharmacy/queue"))["items"]
    assert rx["id"] in [q["id"] for q in queue]
    q_rx = next(q for q in queue if q["id"] == rx["id"])
    assert q_rx["patient"]["allergies"] == "Penicillin"
    done = ok(pharm.post(f"{V}/pharmacy/prescriptions/{rx['id']}/dispense",
                         json={"version": q_rx["version"], "items": [{"prescription_item_id": q_rx["items"][0]["id"],
                                                                      "inventory_item_id": item["id"], "quantity": 10}]}))
    assert done["status"] == "dispensed"
    by_code = ok(pharm.get(f"{V}/inventory/lookup", query_string={"code": item["barcode"]}))
    assert by_code["id"] == item["id"]

    # 5. Lab: dentist requests, lab enters + finalizes, general doctor seeing the patient can read the result.
    glu = ok(mgr.post(f"{V}/lab/tests", json={"code": f"G{uuid.uuid4().hex[:5]}", "name": "Glucose", "unit": "mg/dL",
                                               "ref_low": "70", "ref_high": "110", "price": "20000.00"}))
    req = ok(dent.post(f"{V}/lab/requests", json={"patient_id": pid, "test_ids": [glu["id"]]}))
    req = ok(lab.post(f"{V}/lab/requests/{req['id']}/start", json={"version": req["version"]}))
    req = ok(lab.put(f"{V}/lab/requests/{req['id']}/results",
                     json={"version": req["version"], "items": [{"id": req["items"][0]["id"], "result_value": "145"}]}))
    assert req["items"][0]["abnormal_flag"] == "H"
    req = ok(lab.post(f"{V}/lab/requests/{req['id']}/finalize", json={"version": req["version"]}))
    # The patient now visits General Medicine: found via center-wide lookup and linked (no duplicate profile).
    found = ok(gen.get(f"{V}/patients/lookup", query_string={"q": "Ali Ahmad"}))["items"]
    assert [x["id"] for x in found] == [pid] and found[0]["accessible"] is False
    ok(gen.post(f"{V}/patients/{pid}/link", json={"clinic_id": cl["gen1"]}))
    ok(gen.post(f"{V}/generic/visits", json={"patient_id": pid, "clinic_id": cl["gen1"],
                                              "record": {"chief_complaint": "Fatigue", "vitals": {"bp_systolic": 120}}}))
    assert ok(gen.get(f"{V}/lab/patients/{pid}/requests"))["items"][0]["status"] == "completed"

    # 6. Billing: invoice with discount, partial cash payment, PDF receipt; balance tracked.
    inv = ok(rec.post(f"{V}/billing/invoices", json={"patient_id": pid, "clinic_id": cl["dent1"],
                                                       "items": [{"description": "Filling #14", "qty": 1,
                                                                  "unit_price": "500000.00",
                                                                  "discount_amount": "50000.00"}]}))
    inv = ok(rec.post(f"{V}/billing/invoices/{inv['id']}/issue", json={"version": inv["version"]}))
    pay = ok(rec.post(f"{V}/billing/invoices/{inv['id']}/payments", json={"amount": "300000.00"}))
    inv = ok(rec.get(f"{V}/billing/invoices/{inv['id']}"))
    assert inv["total"] == "450000.00" and inv["paid_total"] == "300000.00" and inv["balance"] == "150000.00"
    assert inv["status"] == "partially_paid"
    assert rec.get(f"{V}/billing/invoices/{inv['id']}/pdf").data[:5] == b"%PDF-"
    assert pay

    # 7. Files: upload to the patient, visible to the clinic, shared explicitly with dermatology, quota tracked.
    up = rec.c.post(f"{V}/files", data={"patient_id": str(pid), "clinic_id": str(cl["dent1"]), "category": "photo",
                                         "files": (io.BytesIO(png_bytes()), "smile.png")},
                    headers=rec._h(None), content_type="multipart/form-data")
    f = ok(up)["items"][0]
    derm = client_for(A["users"]["derm_doc"])
    assert derm.get(f"{V}/files/{f['id']}/content").status_code == 404
    ok(dent.post(f"{V}/files/{f['id']}/shares", json={"target_type": "department",
                                                       "target_id": A["departments"]["dermatology"]}))
    assert derm.get(f"{V}/files/{f['id']}/content").status_code == 200
    assert outsider.get(f"{V}/files/{f['id']}/content").status_code == 404

    # 8. Offline replay: the same X-Op-Id never creates a second record.
    op = {"X-Op-Id": str(uuid.uuid4())}
    body = {"patient_id": pid, "clinic_id": cl["dent1"], "starts_at": f"{DAY}T12:00:00", "ends_at": f"{DAY}T12:30:00"}
    first = rec.post(f"{V}/appointments", json=body, headers=op)
    again = rec.post(f"{V}/appointments", json=body, headers=op)
    assert first.status_code == 201 and again.status_code == 201
    assert again.headers.get("X-Op-Replayed") == "1" and again.get_json()["id"] == first.get_json()["id"]

    # 9. Delete + undo within the window.
    staged = ok(rec.delete(f"{V}/appointments/{first.get_json()['id']}"), 202)
    assert rec.get(f"{V}/appointments/{first.get_json()['id']}").status_code == 404
    ok(rec.post(f"{V}/undo", json={"undo_token": staged["undo_token"]}))
    assert rec.get(f"{V}/appointments/{first.get_json()['id']}").status_code == 200

    # 10. Reports + exports respect scope: manager sees revenue, Excel opens, PDF renders.
    rev = ok(mgr.get(f"{V}/reports/revenue", query_string={"date_from": "2026-01-01", "date_to": "2027-12-31"}))
    assert rev["rows"] is not None
    x = mgr.get(f"{V}/reports/patients/export", query_string={"format": "xlsx", "date_from": "2026-01-01",
                                                              "date_to": "2027-12-31"})
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.data))
    assert wb.active.max_row >= 2
    p = mgr.get(f"{V}/reports/appointments/export", query_string={"format": "pdf", "lang": "ar"})
    assert p.status_code == 200 and p.data[:5] == b"%PDF-"
    assert outsider.get(f"{V}/billing/invoices/{inv['id']}").status_code == 404

    # 11. Complete patient summary: manager sees all departments; dentist doesn't see general-medicine details.
    full = ok(mgr.get(f"{V}/patients/{pid}/summary"))
    scoped = ok(dent.get(f"{V}/patients/{pid}/summary"))
    assert len(str(full)) > len(str(scoped))


def test_new_login_kicks_previous_device(app, world):
    email = world["A"]["users"]["rec_dent"]["email"]
    phone, desk = ApiClient(app), ApiClient(app)
    assert phone.login(email).status_code == 200
    assert desk.login(email).status_code == 200
    assert phone.get(f"{V}/auth/me").get_json()["error"]["code"] == "session_replaced"
