"""Excel / PDF exports: content, column selection, and the same scope/tenant boundaries."""
import io

from openpyxl import load_workbook

from tests.test_reports_helpers import BASE, JAN, q, seed, set_perm


def _export(c, key, fmt="xlsx", **params):
    return c.get(f"{BASE}/{key}/export?{q({**JAN, **params, 'format': fmt})}")


def _sheet_values(resp):
    assert resp.status_code == 200, resp.get_json()
    assert resp.mimetype.endswith("spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def _table(values, first_header):
    """Rows after the header row whose first cell is `first_header`, up to the first empty row."""
    start = next(i for i, r in enumerate(values) if r and r[0] == first_header)
    out = []
    for r in values[start + 1:]:
        if r[0] is None:
            break
        out.append(r)
    return values[start], out


def test_xlsx_manager_content_and_columns(app, world, client_for):
    a = world["A"]
    seed(app, a, names={"P2": "=HYPERLINK(\"http://x\")"})
    m = client_for(a["users"]["manager"])
    r = _export(m, "outstanding")
    assert "attachment" in r.headers["Content-Disposition"] and ".xlsx" in r.headers["Content-Disposition"]
    vals = _sheet_values(r)
    assert vals[0][0] == "Outstanding balances"
    header, rows = _table(vals, "Code")
    assert list(header[:5]) == ["Code", "Patient", "Open invoices", "Balance", "Oldest invoice"]
    assert [row[1] for row in rows[:2]] == ["=HYPERLINK(\"http://x\")", "Patient P1"]  # text, not a formula
    assert float(rows[0][3]) == 80.0 and rows[-1][0] == "Total" and float(rows[-1][3]) == 130.0
    wb = load_workbook(io.BytesIO(r.data))
    cell = next(c for row in wb.active.iter_rows() for c in row if c.value == "=HYPERLINK(\"http://x\")")
    assert cell.data_type == "s"

    # user-selected columns, in the requested order
    vals = _sheet_values(_export(m, "appointments", group_by="clinic", columns="group,no_show,total"))
    header, rows = _table(vals, "Clinic")
    assert list(header) == ["Clinic", "No-show", "Total"]
    assert ("Dental Clinic 1", 1, 3) in [tuple(x) for x in rows]

    # Arabic labels + RTL sheet
    r = _export(m, "patients", lang="ar")
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.active.sheet_view.rightToLeft is True
    assert wb.active["A1"].value == "المرضى الجدد"


def test_department_manager_export_only_department(app, world, client_for):
    a = world["A"]
    seed(app, a)
    h = client_for(a["users"]["dent_head"])
    vals = _sheet_values(_export(h, "clinic_activity"))
    _, rows = _table(vals, "Clinic")
    assert {r[0] for r in rows if r[0] != "Total"} == {"Dental Clinic 1", "Dental Clinic 2"}
    flat = " ".join(str(c) for r in vals for c in r if c is not None)
    assert "Dermatology" not in flat and "Patient P3" not in flat
    vals = _sheet_values(_export(h, "outstanding"))
    flat = " ".join(str(c) for r in vals for c in r if c is not None)
    assert "Patient P1" in flat and "Patient P3" not in flat
    assert _export(h, "patients", department_id=a["departments"]["dermatology"]).status_code == 404


def test_doctor_export_only_clinic(app, world, client_for):
    a = world["A"]
    seed(app, a)
    set_perm(app, a, "dent_doc1", "reports.export", True)
    d = client_for(a["users"]["dent_doc1"])
    vals = _sheet_values(_export(d, "patients", group_by="clinic"))
    _, rows = _table(vals, "Clinic")
    assert [(r[0], r[1]) for r in rows] == [("Dental Clinic 1", 2), ("Total", 2)]
    vals = _sheet_values(_export(d, "outstanding"))
    flat = " ".join(str(c) for r in vals for c in r if c is not None)
    assert "Patient P1" in flat and "Patient P2" not in flat
    assert _export(d, "appointments", clinic_id=a["clinics"]["dent2"]).status_code == 404


def test_cross_tenant_export(app, world, client_for):
    a, b = world["A"], world["B"]
    seed(app, a, names={"P1": "Alpha Secret"})
    seed(app, b, names={"P1": "Bravo One"})
    mb = client_for(b["users"]["manager"])
    for key in ("outstanding", "patients", "clinic_activity"):
        r = _export(mb, key)
        flat = " ".join(str(c) for row in _sheet_values(r) for c in row if c is not None)
        assert "Alpha Secret" not in flat and world["A"]["tag"] not in flat
    vals = _sheet_values(_export(mb, "outstanding"))
    assert "Bravo One" in " ".join(str(c) for r in vals for c in r if c is not None)
    r = _export(mb, "patients", "pdf", clinic_id=a["clinics"]["dent1"])
    assert r.status_code == 404


def test_pdf_exports(app, world, client_for):
    a = world["A"]
    seed(app, a)
    m = client_for(a["users"]["manager"])
    for key, extra in (("revenue", {"group_by": "clinic"}), ("laboratory", {}), ("inventory", {"view": "expiry"}),
                       ("department_activity", {"department_id": a["departments"]["dentistry"]})):
        r = _export(m, key, "pdf", **extra)
        assert r.status_code == 200, (key, r.get_json())
        assert r.mimetype == "application/pdf" and r.data.startswith(b"%PDF")
        assert "attachment" in r.headers["Content-Disposition"]
    r = _export(m, "patients", "pdf", lang="ar", columns="group,new_patients")
    assert r.status_code == 200 and r.data.startswith(b"%PDF")
